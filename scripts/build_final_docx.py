import os
import pypandoc
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def prepare_markdown():
    md_file = r"C:\Users\86449\.gemini\antigravity-ide\brain\4336a428-9dcf-4000-958c-9eb434954937\paper-gemini.md"
    with open(md_file, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # 将图片替换为绝对路径
    md_text = md_text.replace(
        "> **[在此处插入图片：data_distribution.png]**",
        "![各品牌数据量分布直方图](e:/AI Course-Project/results/数据预筛与排序.png)"
    )

    md_text = md_text.replace(
        "> **[在此处插入图片：SimpleCNN_loss_curve.png 与 SimpleCNN_acc_curve.png]**",
        "![SimpleCNN Accuracy](e:/AI Course-Project/results/SimpleCNN_acc_curve.png)\n\n![SimpleCNN Loss](e:/AI Course-Project/results/SimpleCNN_loss_curve.png)\n\n![SimpleCNN 训练过程控制台输出](e:/AI Course-Project/results/训练截图CNN.png)"
    )

    md_text = md_text.replace(
        "> **[在此处插入图片：ResNet50_loss_curve.png 与 ResNet50_acc_curve.png]**",
        "![ResNet50 Accuracy](e:/AI Course-Project/results/ResNet50_acc_curve.png)\n\n![ResNet50 Loss](e:/AI Course-Project/results/ResNet50_loss_curve.png)\n\n![ResNet50 训练过程控制台输出](e:/AI Course-Project/results/训练截图resnet.png)"
    )

    md_text = md_text.replace(
        "> **[在此处插入图片：SimpleCNN_roc_curve.png 与 ResNet50_roc_curve.png]**",
        "![SimpleCNN ROC](e:/AI Course-Project/results/SimpleCNN_roc_curve.png)\n\n![ResNet50 ROC](e:/AI Course-Project/results/ResNet50_roc_curve.png)"
    )

    md_text = md_text.replace(
        "> **[在此处插入图片：ResNet50_confusion_matrix.png]**",
        "![ResNet50 测试集混淆矩阵](e:/AI Course-Project/results/ResNet50_confusion_matrix.png)"
    )

    md_text = md_text.replace(
        "> **[在此处插入图片：ResNet50_error_examples.png]**",
        "![ResNet50 典型的识别错误案例可视化](e:/AI Course-Project/results/ResNet50_error_examples.png)"
    )

    # 增加GUI测试章节
    gui_section = '''
### (六) GUI 图形化界面测试
为了更直观地验证模型的泛化能力与实际应用价值，本文基于 Tkinter 构建了双模型对比测试客户端。可以直接拖入互联网上任意寻找的新图片进行预测。从测试截图可以看出，即使在面临反光、遮挡和角度倾斜的复杂自然场景下，ResNet50 仍能以极高的置信度给出正确预测，而基线 SimpleCNN 往往表现挣扎。

![GUI Test 1](e:/AI Course-Project/results/屏幕截图 2026-06-15 115703.png)

![GUI Test 2](e:/AI Course-Project/results/屏幕截图 2026-06-15 115753.png)

![GUI Test 3](e:/AI Course-Project/results/屏幕截图 2026-06-15 115804.png)
'''
    md_text = md_text.replace("## 7. 结论", gui_section + "\n## 7. 结论")

    temp_md = "e:/AI Course-Project/temp_paper.md"
    with open(temp_md, 'w', encoding='utf-8') as f:
        f.write(md_text)
    return temp_md

def build_pandoc_docx():
    temp_md = prepare_markdown()
    out_docx = "e:/AI Course-Project/《人工智能》课程论文_饮料LOGO分类_终极版.docx"
    
    print("Running pandoc...")
    # 使用 pandoc 将 md 转为 docx
    pypandoc.convert_file(temp_md, 'docx', outputfile=out_docx)
    print("Pandoc generated raw docx.")
    
    print("Post-processing styles...")
    doc = Document(out_docx)
    
    # 强制修改所有段落格式
    for p in doc.paragraphs:
        p.paragraph_format.line_spacing = 1.5
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
            # 根据样式设置字号
            if 'Heading 1' in p.style.name:
                run.font.size = Pt(16)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif 'Heading 2' in p.style.name:
                run.font.size = Pt(15)
            elif 'Heading 3' in p.style.name:
                run.font.size = Pt(14)
            elif 'Caption' in p.style.name or p.text.startswith('图 '):
                run.font.size = Pt(10.5)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if run.font.size is None:
                    run.font.size = Pt(12)
    
    # 强制修改所有表格格式
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10.5)
                        run.font.name = 'Times New Roman'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')

    doc.save(out_docx)
    print(f"Successfully generated ultimate DOCX at {out_docx}")

if __name__ == "__main__":
    build_pandoc_docx()
