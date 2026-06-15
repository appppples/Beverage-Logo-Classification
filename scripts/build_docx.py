import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """Hide cell borders for formula tables"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            edge_el = OxmlElement(f'w:{edge}')
            for key, val in kwargs[edge].items():
                edge_el.set(qn(f'w:{key}'), str(val))
            tcBorders.append(edge_el)
        else:
            edge_el = OxmlElement(f'w:{edge}')
            edge_el.set(qn('w:val'), 'nil')
            tcBorders.append(edge_el)
    tcPr.append(tcBorders)

def add_formula_table(doc, formula_text, formula_num):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(13)
    table.columns[1].width = Cm(2)
    
    cell_formula = table.cell(0, 0)
    p_f = cell_formula.paragraphs[0]
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_f = p_f.add_run(formula_text)
    run_f.font.name = 'Times New Roman'
    run_f.font.size = Pt(12)
    
    cell_num = table.cell(0, 1)
    p_n = cell_num.paragraphs[0]
    p_n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_n = p_n.add_run(formula_num)
    run_n.font.name = 'Times New Roman'
    run_n.font.size = Pt(12)
    
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def create_docx():
    doc = Document()
    
    # 全局样式
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    style.font.size = Pt(12)  # 小四
    style.paragraph_format.line_spacing = 1.5
    
    # 标题样式
    for i in range(1, 4):
        hs = doc.styles[f'Heading {i}']
        hs.font.name = 'SimHei'
        hs._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.bold = True
        
    h1 = doc.styles['Heading 1']
    h1.font.size = Pt(16)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_after = Pt(12)

    h2 = doc.styles['Heading 2']
    h2.font.size = Pt(15)
    
    h3 = doc.styles['Heading 3']
    h3.font.size = Pt(14)

    # 插入标题
    doc.add_heading('基于迁移学习与卷积神经网络的饮料品牌 LOGO 分类系统-gemini', level=1)
    
    md_file = r"C:\Users\86449\.gemini\antigravity-ide\brain\4336a428-9dcf-4000-958c-9eb434954937\paper-gemini.md"
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# 基于'):
            continue  # 跳过大标题
            
        if line.startswith('## '):
            doc.add_heading(line.replace('## ', ''), level=2)
            continue
            
        if line.startswith('### '):
            doc.add_heading(line.replace('### ', ''), level=3)
            continue
            
        # 表格处理
        if line.startswith('|'):
            if '---' in line:
                continue
            row_data = [cell.strip() for cell in line.split('|') if cell.strip()]
            if not in_table:
                in_table = True
                table_data = [row_data]
            else:
                table_data.append(row_data)
            continue
        elif in_table:
            # 渲染表格
            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, r_data in enumerate(table_data):
                for c_idx, c_data in enumerate(r_data):
                    cell = table.cell(r_idx, c_idx)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(c_data)
                    # 表格字体要求：小一号或小半号 (五号 = 10.5pt)
                    run.font.size = Pt(10.5)
                    run.font.name = 'Times New Roman'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
                    if r_idx == 0:
                        run.font.bold = True
            doc.add_paragraph() # 换行
            in_table = False
            table_data = []
            
        # 公式处理
        if line.startswith('$$') and line.endswith('$$'):
            formula = line.replace('$$', '').strip()
            # 提取逗号
            if formula.endswith(','):
                formula = formula[:-1] + " ,"
            add_formula_table(doc, formula, "(4-1)") # 示例编号
            continue
        
        # 图片替换
        if line.startswith('> **[在此处插入图片：'):
            img_name = line.split('：')[1].split(']')[0]
            # 映射到具体图片
            img_map = {
                'data_distribution.png': [r"e:\AI Course-Project\results\数据预筛与排序.png"],
                'SimpleCNN_loss_curve.png 与 SimpleCNN_acc_curve.png': [
                    r"e:\AI Course-Project\results\SimpleCNN_acc_curve.png",
                    r"e:\AI Course-Project\results\SimpleCNN_loss_curve.png",
                    r"e:\AI Course-Project\results\训练截图CNN.png"
                ],
                'ResNet50_loss_curve.png 与 ResNet50_acc_curve.png': [
                    r"e:\AI Course-Project\results\ResNet50_acc_curve.png",
                    r"e:\AI Course-Project\results\ResNet50_loss_curve.png",
                    r"e:\AI Course-Project\results\训练截图resnet.png"
                ],
                'SimpleCNN_roc_curve.png 与 ResNet50_roc_curve.png': [
                    r"e:\AI Course-Project\results\SimpleCNN_roc_curve.png",
                    r"e:\AI Course-Project\results\ResNet50_roc_curve.png"
                ],
                'ResNet50_confusion_matrix.png': [r"e:\AI Course-Project\results\ResNet50_confusion_matrix.png"],
                'ResNet50_error_examples.png': [r"e:\AI Course-Project\results\ResNet50_error_examples.png"]
            }
            if img_name in img_map:
                for img_path in img_map[img_name]:
                    if os.path.exists(img_path):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run = p.add_run()
                        if "训练截图" in img_path:
                            run.add_picture(img_path, width=Inches(6.0))
                        else:
                            run.add_picture(img_path, width=Inches(3.0))
            continue
            
        if line.startswith('> 图') or line.startswith('图 '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line.replace('> ', ''))
            run.font.size = Pt(10.5) # 图注比正文小一号
            continue
            
        if line.startswith('> '):
            continue # 其他提示信息跳过

        # 普通正文
        p = doc.add_paragraph()
        if not line.startswith('表'):
            p.paragraph_format.first_line_indent = Pt(24) # 首行缩进2字符 (12pt * 2)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER # 表题居中
            
        run = p.add_run(line)
        if line.startswith('表'):
            run.font.size = Pt(10.5)

    # 追加 GUI 测试部分
    doc.add_heading('### (六) GUI 图形化界面测试', level=3)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.add_run("为了更直观地验证模型的泛化能力与实际应用价值，本文基于 Tkinter 构建了双模型对比测试客户端。可以直接拖入互联网上任意寻找的新图片进行预测。从测试截图可以看出，即使在面临反光、遮挡和角度倾斜的复杂自然场景下，ResNet50 仍能以极高的置信度给出正确预测，而基线 SimpleCNN 往往表现挣扎。")
    
    gui_imgs = [
        r"e:\AI Course-Project\results\屏幕截图 2026-06-15 115703.png",
        r"e:\AI Course-Project\results\屏幕截图 2026-06-15 115753.png",
        r"e:\AI Course-Project\results\屏幕截图 2026-06-15 115804.png"
    ]
    for img in gui_imgs:
        if os.path.exists(img):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(img, width=Inches(5.0))

    out_path = r"e:\AI Course-Project\《人工智能》课程论文_饮料LOGO分类.docx"
    doc.save(out_path)
    print(f"Successfully generated DOCX at {out_path}")

class Chars(object):
    def __init__(self, value):
        self.value = value
    def __int__(self):
        return int(self.value * 12 * 12700)
    
if __name__ == "__main__":
    create_docx()
