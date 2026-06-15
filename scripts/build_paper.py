"""
论文构建脚本
将 paper_content.md 通过 Pandoc 转换为 Word 文档，
再用 python-docx 进行格式后处理。

用法：python scripts/build_paper.py
"""

import os
import sys
import subprocess
import shutil

# 修复 Windows 终端编码
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAPER_MD = os.path.join(PROJECT_ROOT, "paper_content.md")
TEMPLATE_DOCX = os.path.join(PROJECT_ROOT, "课程论文模板.docx")
OUTPUT_DOCX = os.path.join(PROJECT_ROOT, "《人工智能》课程论文_饮料LOGO分类_v2.docx")
TEMP_DOCX = os.path.join(PROJECT_ROOT, "_temp_pandoc_output.docx")


def find_pandoc():
    """查找 pandoc 可执行文件"""
    # 1. 系统 PATH 中查找
    pandoc = shutil.which("pandoc")
    if pandoc:
        return pandoc
    # 2. 常见安装路径
    candidates = [
        r"C:\Program Files\Pandoc\pandoc.exe",
        r"C:\Users\86449\AppData\Local\Pandoc\pandoc.exe",
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def step1_pandoc_convert():
    """使用 Pandoc 将 Markdown 转换为 docx"""
    pandoc = find_pandoc()
    if not pandoc:
        print("❌ 未找到 Pandoc，尝试用 pypandoc...")
        try:
            import pypandoc
            pypandoc.convert_file(
                PAPER_MD, 'docx',
                outputfile=TEMP_DOCX,
                extra_args=[
                    '--reference-doc', TEMPLATE_DOCX,
                    '--resource-path', PROJECT_ROOT,
                    '--standalone',
                ]
            )
            print("✅ pypandoc 转换完成")
            return True
        except Exception as e:
            print(f"❌ pypandoc 也失败了: {e}")
            return False

    cmd = [
        pandoc,
        PAPER_MD,
        '-o', TEMP_DOCX,
        '--reference-doc', TEMPLATE_DOCX,
        '--resource-path', PROJECT_ROOT,
        '--standalone',
    ]
    print(f"🔧 运行 Pandoc: {' '.join(cmd)}")
    result = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8')
    if result.returncode != 0:
        print(f"❌ Pandoc 出错:\n{result.stderr}")
        return False
    print("✅ Pandoc 转换完成")
    return True


def step2_postprocess():
    """使用 python-docx 进行格式后处理"""
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    print("🔧 后处理格式中...")
    doc = Document(TEMP_DOCX)

    # 1. 设置页面尺寸和边距 (A4, 上下2.5cm 左3cm 右2cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # 2. 修正所有段落的字体
    for para in doc.paragraphs:
        # 设置行距 1.5 倍
        para.paragraph_format.line_spacing = 1.5

        style_name = para.style.name if para.style else ""

        for run in para.runs:
            # 西文字体设为 Times New Roman
            run.font.name = 'Times New Roman'
            # 中文字体设为宋体
            rpr = run._element.get_or_add_rPr()
            rFonts = rpr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = run._element.makeelement(qn('w:rFonts'), {})
                rpr.insert(0, rFonts)
            rFonts.set(qn('w:eastAsia'), '宋体')

            # 根据样式设置字号
            if 'Heading 1' in style_name:
                run.font.size = Pt(16)  # 三号
                run.font.bold = True
                # 标题字体用黑体
                rFonts.set(qn('w:eastAsia'), '黑体')
            elif 'Heading 2' in style_name:
                run.font.size = Pt(15)  # 小三
                run.font.bold = True
                rFonts.set(qn('w:eastAsia'), '黑体')
            elif 'Heading 3' in style_name:
                run.font.size = Pt(14)  # 四号
                run.font.bold = True
                rFonts.set(qn('w:eastAsia'), '黑体')
            elif 'Heading 4' in style_name:
                run.font.size = Pt(13)  # 小四偏大
                run.font.bold = True
                rFonts.set(qn('w:eastAsia'), '黑体')
            elif 'Caption' in style_name or para.text.startswith('图'):
                run.font.size = Pt(10.5)  # 五号（图注）
            else:
                if run.font.size is None:
                    run.font.size = Pt(12)  # 小四（正文）

        # Heading 1 居中
        if 'Heading 1' in style_name:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 图注居中
        if para.text.startswith('图') and len(para.text) < 80:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 正文首行缩进（排除标题、图注、列表等）
        if ('Heading' not in style_name
            and 'Caption' not in style_name
            and 'List' not in style_name
            and not para.text.startswith('图')
            and not para.text.startswith('|')
            and not para.text.startswith('[')
            and not para.text.startswith('**[')
            and para.text.strip()
            and not para.text.startswith('>')
        ):
            para.paragraph_format.first_line_indent = Pt(24)  # 2字符

    # 3. 修正表格格式
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(10.5)  # 表格字体五号
                        run.font.name = 'Times New Roman'
                        rpr = run._element.get_or_add_rPr()
                        rFonts = rpr.find(qn('w:rFonts'))
                        if rFonts is None:
                            rFonts = run._element.makeelement(qn('w:rFonts'), {})
                            rpr.insert(0, rFonts)
                        rFonts.set(qn('w:eastAsia'), '宋体')

    # 4. 保存
    doc.save(OUTPUT_DOCX)
    print(f"✅ 格式后处理完成")


def step3_cleanup():
    """清理临时文件"""
    if os.path.exists(TEMP_DOCX):
        os.remove(TEMP_DOCX)
        print("🧹 临时文件已清理")


def main():
    print("=" * 60)
    print("📝 课程论文构建脚本")
    print("=" * 60)

    # 检查输入文件
    if not os.path.exists(PAPER_MD):
        print(f"❌ 找不到论文内容文件: {PAPER_MD}")
        return

    print(f"  📄 输入: {PAPER_MD}")
    print(f"  📋 模板: {TEMPLATE_DOCX}")
    print(f"  📦 输出: {OUTPUT_DOCX}")

    # Step 1: Pandoc 转换
    print("\n" + "-" * 40)
    print("Step 1: Pandoc Markdown → DOCX")
    print("-" * 40)
    if not step1_pandoc_convert():
        print("❌ 转换失败，请检查 Pandoc 安装")
        return

    # Step 2: 格式后处理
    print("\n" + "-" * 40)
    print("Step 2: python-docx 格式后处理")
    print("-" * 40)
    step2_postprocess()

    # Step 3: 清理
    step3_cleanup()

    # 检查输出
    if os.path.exists(OUTPUT_DOCX):
        size_mb = os.path.getsize(OUTPUT_DOCX) / (1024 * 1024)
        print(f"\n{'=' * 60}")
        print(f"🎉 论文生成成功!")
        print(f"   📦 文件: {OUTPUT_DOCX}")
        print(f"   📏 大小: {size_mb:.1f} MB")
        print(f"\n   ⚠️  请在 Word 中打开检查格式，并手动添加封面页")
        print(f"   ⚠️  理论部分的占位图需要自行用 PPT/AI 替换")
        print(f"{'=' * 60}")
    else:
        print("❌ 输出文件未生成")


if __name__ == "__main__":
    main()
